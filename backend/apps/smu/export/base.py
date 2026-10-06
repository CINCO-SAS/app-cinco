"""Estilos y helpers compartidos por las plantillas de exporte a Excel.

Puerto de ``plantillaEmergencia.py`` (y de sus hermanas ``plantillaPreventivoAA.py``
y ``plantillaCorrectivos.py``): misma paleta, mismos bordes y misma maquetación
de 12 columnas (A..L), pero con las funciones parametrizadas por ``ws`` para
poder construir varias hojas en el mismo proceso.

El objetivo es que el archivo que descarga el botón "Exportar" de la tabla de
gestión se vea idéntico a la plantilla que ya conocen en Excel.
"""
from __future__ import annotations

import math
import re
from typing import NamedTuple

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .imagenes import insertar_foto_en_celda, ruta_absoluta

# ── Rellenos de color ────────────────────────────────────────────────────────
FILL_BURGUNDY = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
FILL_YELLOW = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
FILL_LIGHT_YELLOW = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
FILL_PINK_MAGENTA = PatternFill(start_color="FF0066", end_color="FF0066", fill_type="solid")
FILL_LIGHT_GRAY = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
FILL_GRIS_MEDIO = PatternFill(start_color="A6A6A6", end_color="A6A6A6", fill_type="solid")
FILL_GRAY = PatternFill(start_color="808080", end_color="808080", fill_type="solid")
FILL_WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_ROSA = PatternFill(start_color="FBE3D5", end_color="FBE3D5", fill_type="solid")

# ── Fuentes (Calibri 9pt) ───────────────────────────────────────────────────
FONT_HEADER_WHITE = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
FONT_HEADER_BLACK = Font(name="Calibri", size=9, bold=True, color="000000")
FONT_DATA_BOLD = Font(name="Calibri", size=9, bold=True, color="000000")
FONT_DATA_REGULAR = Font(name="Calibri", size=9, bold=False, color="000000")

# ── Alineación ──────────────────────────────────────────────────────────────
ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

# ── Bordes ──────────────────────────────────────────────────────────────────
SIDE_THIN = Side(style="thin", color="BFBFBF")
SIDE_BLANCO = Side(style="medium", color="FFFFFF")
SIDE_NEGRO = Side(style="thick", color="000000")
SIDE_TABLA = Side(style="medium", color="000000")

BORDER_THIN = Border(left=SIDE_THIN, right=SIDE_THIN, top=SIDE_THIN, bottom=SIDE_THIN)
BORDER_TABLA = Border(left=SIDE_TABLA, right=SIDE_TABLA, top=SIDE_TABLA, bottom=SIDE_TABLA)
BORDER_BLANCO = Border(left=SIDE_NEGRO, right=SIDE_NEGRO, top=SIDE_BLANCO, bottom=SIDE_BLANCO)
BORDER_NEGRO = Border(left=SIDE_NEGRO, right=SIDE_NEGRO, top=SIDE_NEGRO, bottom=SIDE_NEGRO)
BORDER_SPACE = Border(left=SIDE_NEGRO, right=SIDE_NEGRO, top=SIDE_NEGRO, bottom=SIDE_NEGRO)

# Documento de 12 columnas: A..L
COL_INICIO = 1
COL_FIN = 12

COLUMN_WIDTHS = {
    "A": 12.0, "B": 12.0, "C": 12.0, "D": 13.0, "E": 10.0, "F": 12.0,
    "G": 14.0, "H": 14.0, "I": 13.0, "J": 13.0, "K": 15.0, "L": 15.0,
}


def aplicar_anchos(ws, columnas: dict[str, float] | None = None) -> None:
    """Asigna el ancho de cada columna (A..L) de la hoja."""
    for col, ancho in (columnas or COLUMN_WIDTHS).items():
        ws.column_dimensions[col].width = ancho


def fila_banner(
    ws,
    fila: int,
    texto: str,
    fill: PatternFill,
    font: Font,
    border: Border = BORDER_BLANCO,
    col_ini: int = COL_INICIO,
    col_fin: int = COL_FIN,
) -> None:
    """Fila tipo 'banner': una sola celda fusionada con color de fondo."""
    ws.merge_cells(start_row=fila, start_column=col_ini, end_row=fila, end_column=col_fin)
    for col in range(col_ini, col_fin + 1):
        celda = ws.cell(row=fila, column=col)
        celda.fill = fill
        celda.font = font
        celda.border = border
        celda.alignment = ALIGN_CENTER
    ws.cell(row=fila, column=col_ini).value = texto


def fila_espaciadora(ws, fila: int, alto: int = 6) -> None:
    """Fila delgada, usada solo para separar secciones."""
    ws.row_dimensions[fila].height = alto
    ws.merge_cells(start_row=fila, start_column=COL_INICIO, end_row=fila, end_column=COL_FIN)
    ws.cell(row=fila, column=COL_INICIO).border = BORDER_SPACE


def fusionar(
    ws,
    fila_ini: int,
    col_ini: int,
    fila_fin: int,
    col_fin: int,
    valor=None,
    fill: PatternFill | None = None,
    font: Font | None = None,
    border: Border | None = None,
    align: Alignment = ALIGN_CENTER,
):
    """
    Fusiona un rango de celdas y le aplica estilo a TODO el bloque.

    Los bordes se aplican solo en las aristas exteriores del rango para evitar
    el efecto de "recorte" visual en celdas fusionadas.
    """
    ws.merge_cells(start_row=fila_ini, start_column=col_ini, end_row=fila_fin, end_column=col_fin)

    for f in range(fila_ini, fila_fin + 1):
        for c in range(col_ini, col_fin + 1):
            celda = ws.cell(row=f, column=c)
            if fill:
                celda.fill = fill
            if font:
                celda.font = font
            if border:
                lado_izq = border.left if c == col_ini else Side(style=None)
                lado_der = border.right if c == col_fin else Side(style=None)
                lado_arriba = border.top if f == fila_ini else Side(style=None)
                lado_abajo = border.bottom if f == fila_fin else Side(style=None)
                celda.border = Border(
                    left=lado_izq, right=lado_der,
                    top=lado_arriba, bottom=lado_abajo,
                )
            celda.alignment = align

    if valor is not None:
        ws.cell(row=fila_ini, column=col_ini).value = valor

    # Dirección de la celda principal del bloque (ancla para fotos u objetos).
    return ws.cell(row=fila_ini, column=col_ini).coordinate


def fila_encabezados_tabla(
    ws,
    fila: int,
    encabezados: list[str],
    col_ini: int = 1,
    fill: PatternFill = FILL_ROSA,
) -> None:
    """Escribe los títulos de columna de una tabla (fila rosada/gris)."""
    for i, texto in enumerate(encabezados):
        celda = ws.cell(row=fila, column=col_ini + i)
        celda.value = texto
        celda.fill = fill
        celda.font = FONT_HEADER_BLACK
        celda.alignment = ALIGN_CENTER
        celda.border = BORDER_NEGRO


def estilizar_fila(
    ws,
    fila: int,
    fill: PatternFill | None = None,
    font: Font | None = None,
    border: Border | None = None,
    alto: int | None = None,
    col_ini: int = COL_INICIO,
    col_fin: int = COL_FIN,
    align: Alignment = ALIGN_CENTER,
) -> None:
    """Aplica estilos a una fila completa (o rango de columnas) sin fusionar."""
    if alto:
        ws.row_dimensions[fila].height = alto
    for col in range(col_ini, col_fin + 1):
        celda = ws.cell(row=fila, column=col)
        if fill:
            celda.fill = fill
        if font:
            celda.font = font
        if border:
            celda.border = border
        if align:
            celda.alignment = align


def texto_o_vacio(valor) -> str:
    """Normaliza valores nulos/vacíos a cadena vacía (el Excel no debe mostrar None)."""
    if valor is None:
        return ""
    return str(valor)


def formato_fecha(valor, con_hora: bool = False) -> str:
    """`datetime`/`date` → ``30/09/2026`` (o ``30/09/2026 10:12``); vacío → ``""``."""
    if not valor:
        return ""
    try:
        fecha = valor.date() if hasattr(valor, "date") and not isinstance(valor, str) else valor
        if hasattr(fecha, "strftime"):
            if con_hora and hasattr(valor, "strftime"):
                return valor.strftime("%d/%m/%Y %H:%M")
            return fecha.strftime("%d/%m/%Y")
        return str(valor)
    except (ValueError, TypeError):
        return str(valor)


def numero(valor, mostrar_cero: bool = False) -> str | None:
    """`Decimal`/número → ``12.34`` sin ceros sobrantes; 0 (default) → vacío.

    Los modelos traen ``default=0.00`` en casi todos los campos numéricos, así
    que un cero no dice nada: se muestra vacío salvo que se pida lo contrario.
    """
    if valor is None:
        return None
    try:
        flotante = float(valor)
    except (ValueError, TypeError):
        return texto_o_vacio(valor) or None
    if flotante == 0 and not mostrar_cero:
        return None
    return f"{flotante:g}"


# ── Composición de secciones (compartido por todos los generadores) ─────────
# Anchos de los bloques de texto (en caracteres) para estimar cuántas filas
# necesita cada párrafo. A..L ≈ 155; A..F ≈ 71; G..L ≈ 84.
ANCHO_COMPLETO = 145
ANCHO_MITAD = 68
ALTO_FILA = 15  # puntos: una línea de Calibri 9


def filas_para_texto(texto: str, chars_por_fila: int = ANCHO_COMPLETO) -> int:
    """Cantidad de filas que ocupa un texto (mínimo 2, como la plantilla)."""
    if not texto:
        return 2
    lineas = math.ceil(len(texto) / chars_por_fila)
    # Considera los saltos de línea reales del texto.
    lineas += texto.count("\n")
    return max(2, lineas)


def bloque_texto(ws, fila: int, texto: str, col_ini: int, col_fin: int) -> int:
    """Pinta un área de texto gris con altura según el contenido. Devuelve la próxima fila."""
    filas = filas_para_texto(texto, ANCHO_COMPLETO if col_ini == 1 else ANCHO_MITAD)
    fusionar(
        ws, fila, col_ini, fila + filas - 1, col_fin,
        valor=texto or None,
        fill=FILL_LIGHT_GRAY, font=FONT_DATA_REGULAR,
        border=BORDER_TABLA, align=ALIGN_LEFT,
    )
    for r in range(fila, fila + filas):
        ws.row_dimensions[r].height = ALTO_FILA
    return fila + filas


def par_de_campos(ws, fila: int, izquierda, derecha) -> None:
    """Fila 'etiqueta | valor | etiqueta | valor' de dos columnas de datos."""
    etiqueta_1, valor_1 = izquierda
    etiqueta_2, valor_2 = derecha

    fusionar(ws, fila, 1, fila, 3, valor=etiqueta_1, font=FONT_DATA_BOLD,
             border=BORDER_TABLA, align=ALIGN_LEFT)
    fusionar(ws, fila, 4, fila, 6, valor=valor_1 or None, fill=FILL_YELLOW,
             font=FONT_DATA_REGULAR, border=BORDER_TABLA, align=ALIGN_LEFT)
    fusionar(ws, fila, 7, fila, 9, valor=etiqueta_2, font=FONT_DATA_BOLD,
             border=BORDER_TABLA, align=ALIGN_LEFT)
    fusionar(ws, fila, 10, fila, 12, valor=valor_2 or None, fill=FILL_YELLOW,
             font=FONT_DATA_REGULAR, border=BORDER_TABLA, align=ALIGN_LEFT)


def pares_con_datos(
    pares: list[tuple[tuple[str, str], tuple[str, str]]],
    primero_siempre: bool = True,
):
    """Filtra las filas donde ambos valores están vacíos."""
    resultado = []
    for indice, par in enumerate(pares):
        (_, valor_1), (_, valor_2) = par
        if indice == 0 and primero_siempre:
            resultado.append(par)
        elif texto_o_vacio(valor_1).strip() or texto_o_vacio(valor_2).strip():
            resultado.append(par)
    return resultado


def sub_titulo(ws, fila: int, texto: str) -> int:
    """Banner gris de sub-sección (ej. 'DATOS GENERALES AIRES ACONDICIONADOS')."""
    fusionar(ws, fila, COL_INICIO, fila, COL_FIN, valor=texto,
             fill=FILL_GRAY, font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
    return fila + 1


def encabezados_grupos(
    ws,
    fila: int,
    grupos: list[tuple[int, int]],
    encabezados: list[str],
    fill: PatternFill = FILL_GRIS_MEDIO,
    font: Font = FONT_HEADER_WHITE,
) -> int:
    """Fila de encabezados con columnas fusionadas (grupos de A..L)."""
    for (ini, fin), texto in zip(grupos, encabezados):
        fusionar(ws, fila, ini, fila, fin, valor=texto, fill=fill,
                 font=font, border=BORDER_NEGRO)
    return fila + 1


def datos_grupos(
    ws,
    fila: int,
    grupos: list[tuple[int, int]],
    valores,
    fill: PatternFill = FILL_YELLOW,
    align: Alignment = ALIGN_LEFT,
) -> int:
    """Fila de datos con las mismas fusiones que sus encabezados."""
    for (ini, fin), valor in zip(grupos, valores):
        fusionar(ws, fila, ini, fila, fin, valor=valor or None, fill=fill,
                 font=FONT_DATA_REGULAR, border=BORDER_TABLA, align=align)
    return fila + 1


def fila_con_etiqueta(
    ws,
    fila: int,
    etiqueta: str,
    n_columnas: int,
    col_datos_ini: int = 2,
    fill_etiqueta: PatternFill = FILL_ROSA,
) -> int:
    """Fila 'AA. 1' / 'COMPRESOR 1': etiqueta rosa en A + celdas amarillas."""
    celda = ws.cell(row=fila, column=1)
    celda.value = etiqueta
    celda.fill = fill_etiqueta
    celda.font = FONT_DATA_BOLD
    celda.alignment = ALIGN_CENTER
    celda.border = BORDER_TABLA

    for col in range(col_datos_ini, col_datos_ini + n_columnas):
        celda = ws.cell(row=fila, column=col)
        celda.fill = FILL_YELLOW
        celda.font = FONT_DATA_REGULAR
        celda.alignment = ALIGN_CENTER
        celda.border = BORDER_TABLA
    return fila + 1


# Cajas de foto: 3 por fila, 4 columnas cada una (A-D, E-H, I-L). El formato
# tiene 12 columnas; con dos cajas de 6 la imagen (260 px de máximo) dejaba
# medio hueco sin usar en cada fila.
COLUMNAS_CAJA_FOTO = ((1, 4), (5, 8), (9, 12))

# Aproximadamente cuántos caracteres caben en una caja de 4 columnas: se usa
# de más para nunca recortar un rótulo largo.
_CARACTERES_POR_CAJA = 40


def _alto_rotulo(rotulos) -> float:
    """Alta de la fila del rótulo: un título largo ocupa dos líneas."""
    lineas = max(
        1,
        max(math.ceil(len(rotulo) / _CARACTERES_POR_CAJA) for rotulo in rotulos),
    )
    return 15 * lineas + 5


def bloque_fotos(ws, fila: int, evidencias, titulo=None) -> int:
    """Tres fotos por fila (cajas de 4 columnas: A-D, E-H, I-L).

    ``titulo`` recibe la evidencia y devuelve el rótulo de su caja; por defecto
    usa el orden de la evidencia (como el formato de emergencias).

    Solo se dibuja una caja por foto que exista: si en la última fila sobra
    un hueco se queda sin celdas, no con una caja gris vacía.
    """
    if titulo is None:
        titulo = lambda e: f"FOTO EVIDENCIA {e.orden}"  # noqa: E731

    fotos = list(evidencias)
    ancho = len(COLUMNAS_CAJA_FOTO)
    grupos = [fotos[i:i + ancho] for i in range(0, len(fotos), ancho)]

    for grupo in grupos:
        rotulos = [titulo(evidencia) for evidencia in grupo]
        for lado, rotulo in enumerate(rotulos):
            col_ini, col_fin = COLUMNAS_CAJA_FOTO[lado]
            fusionar(ws, fila, col_ini, fila, col_fin, valor=rotulo,
                     fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE,
                     border=BORDER_NEGRO)
        ws.row_dimensions[fila].height = _alto_rotulo(rotulos)
        fila += 1

        anclas = []
        for lado in range(len(grupo)):
            col_ini, col_fin = COLUMNAS_CAJA_FOTO[lado]
            anclas.append(fusionar(ws, fila, col_ini, fila + 5, col_fin,
                                   fill=FILL_LIGHT_GRAY, font=FONT_DATA_REGULAR,
                                   border=BORDER_TABLA))
        for r in range(fila, fila + 6):
            ws.row_dimensions[r].height = 25

        for lado, evidencia in enumerate(grupo):
            insertar_foto_en_celda(ws, evidencia.ruta_archivo, anclas[lado])

        fila += 6
        fila_espaciadora(ws, fila)
        fila += 1
    return fila


class Numerador:
    """Asigna el número de sección solo a las que se van a pintar."""

    def __init__(self):
        self.valor = 0

    def titulo(self, texto: str) -> str:
        self.valor += 1
        return f"{self.valor}. {texto}"


# Identificador interno de una fila de evidencia (``ev-1790875901691``):
# sirve para emparejar en el backend, pero no significa nada para quien lee
# el Excel, así que nunca se usa como cabecera de foto.
_ORIGEN_INTERNO = re.compile(r"^[A-Za-z]{1,8}-\d{6,}$")


def rotulo_evidencia(evidencia) -> str:
    """Cabecera de la caja de una foto, en este orden:

    1. La **descripción** que el usuario escribió para esa imagen
       (``FotoSlotWithDescripcion`` → ``evidencias.descripcion``): es la única
       que dice en lenguaje humano qué se ve dentro de la foto, y el propio
       formulario pide que "cada imagen incluya una descripción".
    2. Si no escribió nada, el nombre del campo con espacios
       (``condensadora_pos_mantenimiento`` → ``CONDENSADORA POS MANTENIMIENTO``).
    3. Si ese nombre es un identificador interno sin lectura humana
       (``ev-1790875901691``), el orden de la foto.
    """
    descripcion = " ".join(
        texto_o_vacio(getattr(evidencia, "descripcion", "")).split()
    )
    if descripcion:
        return descripcion.upper()

    origen = texto_o_vacio(getattr(evidencia, "campo_origen", "")).strip()
    if not origen or _ORIGEN_INTERNO.match(origen):
        return f"FOTO EVIDENCIA {getattr(evidencia, 'orden', 1)}"
    limpio = " ".join(origen.replace("_", " ").split())
    return limpio.upper()


class AdjuntoFotografico(NamedTuple):
    """Imagen del formulario que **no** es una ``evidencia``."""

    orden: int
    ruta_archivo: str
    etiqueta: str


def imagenes_adjuntas(actividad) -> list[AdjuntoFotografico]:
    """Todas las imágenes del formulario que ningún exporte pintaba.

    Cubren los campos de imagen del modelo que quedaban fuera del Excel:

    - ``certificados.ruta_certificado`` / ``ruta_foto_sitio`` (certificado
      laboral y su foto en sitio);
    - ``tecnicos.ruta_foto`` / ``ruta_certificado`` / ``ruta_firma``;
    - ``firmas_cierre.ruta_firma``.

    Solo devuelve rutas que **existen en disco**: un campo con texto libre
    (una firma escrita a mano, un ``None``) no debe abrir una caja vacía en el
    reporte. Si dos campos apuntan al mismo archivo se pinta una sola vez.
    """
    items: list[AdjuntoFotografico] = []
    vistos: set[str] = set()

    def agregar(ruta, etiqueta: str) -> None:
        limpia = texto_o_vacio(ruta).strip()
        if not limpia or limpia in vistos:
            return
        destino = ruta_absoluta(limpia)
        if destino is None or not destino.is_file():
            return
        vistos.add(limpia)
        # Los rótulos de caja van en mayúsculas, igual que el resto de
        # cabeceras del reporte.
        items.append(
            AdjuntoFotografico(len(items) + 1, limpia, etiqueta.strip().upper())
        )

    for certificado in actividad.certificados.all():
        detalle = " · ".join(
            parte for parte in (
                texto_o_vacio(certificado.tipo).strip(),
                texto_o_vacio(certificado.categoria).strip(),
            ) if parte
        )
        sufijo = f" — {detalle}" if detalle else ""
        agregar(certificado.ruta_certificado, f"CERTIFICADO LABORAL{sufijo}")
        agregar(certificado.ruta_foto_sitio,
                f"FOTO DEL CERTIFICADO EN SITIO{sufijo}")

    for tecnico in actividad.tecnicos.all():
        nombre = texto_o_vacio(tecnico.nombre).strip()
        sufijo = f" — {nombre}" if nombre else ""
        agregar(tecnico.ruta_foto, f"FOTO DEL TÉCNICO{sufijo}")
        agregar(tecnico.ruta_certificado, f"CERTIFICADO DEL TÉCNICO{sufijo}")
        agregar(tecnico.ruta_firma, f"FIRMA DEL TÉCNICO{sufijo}")

    for firma in actividad.firmas_cierre.all():
        partes = [
            parte for parte in (
                texto_o_vacio(firma.rol).strip().upper(),
                texto_o_vacio(firma.nombre).strip(),
            ) if parte
        ]
        rotulo = f"FIRMA {' · '.join(partes)}" if partes else "FIRMA"
        agregar(firma.ruta_firma, rotulo)

    return items


def bloque_imagenes_adjuntas(ws, fila: int, actividad, titulo) -> int:
    """Sección final con certificados, firmas y fotos de técnicos.

    Devuelve ``fila`` sin tocar nada si no hay ninguna imagen que pintar, de
    modo que un reporte sin certificados no gana una sección en blanco.
    ``titulo`` puede ser un callable: se invoca solo cuando sí hay imágenes,
    para que los contadores de sección numerados (``Numerador``) no salten.
    """
    items = imagenes_adjuntas(actividad)
    if not items:
        return fila

    rotulo = titulo() if callable(titulo) else titulo
    fila_banner(ws, fila, rotulo, FILL_BURGUNDY, FONT_HEADER_WHITE)
    fila += 1
    fila_espaciadora(ws, fila)
    fila += 1
    return bloque_fotos(ws, fila, items, titulo=lambda item: item.etiqueta)


def descripcion_material(material) -> str | None:
    """Descripción del material con su código SAP, si lo tiene.

    La plantilla tiene **una sola columna** de descripción, así que se unen
    los dos campos en vez de quedarse solo con el código.
    """
    descripcion = texto_o_vacio(material.descripcion).strip()
    codigo = texto_o_vacio(material.texto_sap).strip()
    if descripcion and codigo:
        return f"{descripcion} (SAP {codigo})"
    return descripcion or codigo or None


def responsable(actividad) -> str:
    nombre = texto_o_vacio(actividad.responsable_nombre).strip()
    cedula = texto_o_vacio(actividad.responsable_cedula).strip()
    if nombre and cedula:
        return f"{nombre} (CC {cedula})"
    return nombre or (f"CC {cedula}" if cedula else "")
