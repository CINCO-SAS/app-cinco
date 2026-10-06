"""Exporte del preventivo de AIRE ACONDICIONADO a Excel.

Port dinámico de ``plantillaPreventivoAA.py``: mismos banners (plan de
mantenimiento preventivo), misma tabla de datos generales con la columna de
etiqueta en A, las mini tablas de temperaturas/condensadora y la lista de
chequeo de 39 parámetros, pero con las filas que pide la ficha real:

- una fila de equipo por dato no vacío (si no hay detalle, no existe la fila),
- una fila por cada compresor / condensadora / manejadora registrada,
- solo los ítems del checklist que el usuario diligenció,
- una fila de acciones por cada unidad AA (``aa_acciones.orden``),
- anexos fotográficos al final (el formato original no traía fotos, pero el
  formulario sí las guarda y el usuario las exige en el archivo).
"""
from __future__ import annotations

from openpyxl import Workbook

from apps.smu.models import SmuActividad, SmuDetallePreventivoAa

from . import base
from .base import (
    BORDER_NEGRO,
    BORDER_TABLA,
    FILL_BURGUNDY,
    FILL_GRIS_MEDIO,
    FILL_ROSA,
    FONT_DATA_BOLD,
    FONT_DATA_REGULAR,
    FONT_HEADER_BLACK,
    FONT_HEADER_WHITE,
    bloque_fotos,
    bloque_texto,
    encabezados_grupos,
    estilizar_fila,
    fila_banner,
    fila_con_etiqueta,
    fila_encabezados_tabla,
    fusionar,
    par_de_campos,
    pares_con_datos,
    sub_titulo,
    texto_o_vacio,
)
from .checklist_aa import parametro_de

HOJA = "Preventivo AA"

# ── Tablas de la plantilla (columnas B..L, etiqueta de fila en A) ────────────
ENCABEZADOS_EQUIPO = [
    "MARCA", "MODELO", "SERIAL", "ESTADO DEL EQUIPO", "TIPO AIRE",
    "ALIMENTACIÓN AC", "VOLTAJE ENTRADA (VAC)", "CORRIENTE (AMP)",
    "CAPACIDAD BTU", "GESTIÓN REMOTA-IP", "CANT. COMPRESORES",
]  # 11 columnas de datos → B..L

ENCABEZADOS_TEMPERATURAS = [
    "TEMP. CUARTO EQUIPO", "TEMP. AA ENTRADA", "TEMP. AA SALIDA",
    "TEMP. DISPLAY", "TEMP. TERMOSTATO", "AJUSTE TERMOSTATO",
    "TEMP. TERMOSTATO POS AJUSTE",
]  # 7 grupos sin etiqueta → A..L
GRUPOS_7 = [(1, 2), (3, 4), (5, 6), (7, 8), (9, 10), (11, 11), (12, 12)]

ENCABEZADOS_COMPRESOR = [
    "MARCA", "SERIAL", "TIPO REFRIGERANTE", "MODELO",
    "AISLAMIENTO ELÉCTRICO (MΩ)", "PRESIÓN SUCCIÓN", "PRESIÓN DESCARGA",
    "NIVEL DE ACEITE", "VL1/VL2/VL3", "AMPL1/AMPL2/AMPL3",
]  # 10 columnas → B..K

ENCABEZADOS_CONDENSADORA = [
    "MARCA", "MODELO", "SERIAL", "TEMPERATURA ENTRADA",
    "TEMPERATURA SALIDA", "DIÁMETRO EJE", "DIÁMETRO ASPAS",
]  # etiqueta en A + 7 grupos → B..L
GRUPOS_CONDENSADORA = [
    (2, 3), (4, 5), (6, 7), (8, 9), (10, 10), (11, 11), (12, 12),
]

ENCABEZADOS_MANEJADORA = [
    "MARCA", "MODELO", "TIPO", "TIPO DE FILTRO", "TIPO CORREA",
    "MARCA MOTOR", "ALIMENTACIÓN AC", "VOLTAJE DEL MOTOR",
    "CORRIENTE MOTOR", "AISLAMIENTO (MΩ)", "SERIAL MOTOR",
]  # 11 columnas → B..L

# Claves tri-estado en el mismo orden que CLAVES_ACCIONES_AA (frontend) y con
# los títulos de la plantilla.
ACCIONES_AA = [
    ("limpieza_serpentines", "LIMPIEZA DE SERPENTINES"),
    ("ajuste_elementos_control", "AJUSTE ELEMENTOS DE CONTROL"),
    ("adicion_refrigerante", "ADICIÓN REFRIGERANTE"),
    ("estado_drenajes", "ESTADO DRENAJES"),
    ("lubricacion_componentes", "LUBRICACIÓN DE COMPONENTES"),
    ("cambio_filtros_secado", "CAMBIO FILTROS SECADO"),
    ("alineacion_poleas", "ALINEACIÓN DE POLEAS"),
    ("cambio_componentes_electronicos", "CAMBIO COMPONENTES ELECTRÓNICOS"),
    ("cambio_compresor", "CAMBIO DE COMPRESOR"),
    ("cambio_correas_filtros", "CAMBIO DE CORREAS/FILTROS"),
    ("otras_reparaciones", "OTRAS REPARACIONES"),
]

# Traducción de los valores de SmuPreventivoAaAccion.VALOR_CHOICES.
VALOR_LEGIBLE = {
    "SI": "SÍ", "NO": "NO", "NA": "N/A",
    "BUENO": "Bueno", "REGULAR": "Regular", "MALO": "Malo",
}


# ── Datos de la actividad ────────────────────────────────────────────────────

def _detalle(actividad: SmuActividad) -> SmuDetallePreventivoAa | None:
    try:
        return actividad.detalle_preventivo_aa
    except SmuDetallePreventivoAa.DoesNotExist:
        return None


def _tecnicos(actividad: SmuActividad) -> list[str]:
    return [
        f"{t.nombre} (CC {t.cedula})" if t.cedula else t.nombre
        for t in actividad.tecnicos.all()
    ]


def _pares_informacion_general(actividad: SmuActividad):
    tecnicos = _tecnicos(actividad)
    izquierda = [
        ("NOMBRE DE ESTACIÓN:", texto_o_vacio(actividad.nombre_estacion)),
        ("DEPARTAMENTO:", texto_o_vacio(actividad.departamento)),
        ("DIRECCIÓN:", texto_o_vacio(actividad.direccion)),
        ("TIPO DE ESTACIÓN:", texto_o_vacio(actividad.tipo_estacion)),
        ("SITE OWNER:", texto_o_vacio(actividad.site_owner)),
        ("EMPRESA:", texto_o_vacio(actividad.empresa)),
    ]
    derecha = [
        ("RESPONSABLE 1:", tecnicos[0] if tecnicos else base.responsable(actividad)),
        ("RESPONSABLE 2:", tecnicos[1] if len(tecnicos) > 1 else ""),
        ("OT:", texto_o_vacio(actividad.codigo_ot)),
        ("FECHA INICIO EJECUCIÓN:", base.formato_fecha(actividad.fecha_inicio, con_hora=True)),
        ("FECHA FIN ACTIVIDAD:", base.formato_fecha(actividad.fecha_fin, con_hora=True)),
        ("CATEGORÍA:", actividad.get_categoria_display()),
    ]
    return list(zip(izquierda, derecha))


def _datos_equipo(detalle: SmuDetallePreventivoAa | None) -> list:
    """Valores de la fila 'AA. 1' en el orden de ``ENCABEZADOS_EQUIPO``."""
    if detalle is None:
        return []
    return [
        texto_o_vacio(detalle.marca),
        texto_o_vacio(detalle.modelo),
        texto_o_vacio(detalle.serial),
        texto_o_vacio(detalle.estado_equipo),
        texto_o_vacio(detalle.tipo_aire),
        texto_o_vacio(detalle.alimentacion_ac),
        base.numero(detalle.voltaje_entrada),
        base.numero(detalle.corriente),
        base.numero(detalle.capacidad_btu),
        texto_o_vacio(detalle.gestion_remota_ip),
        base.numero(detalle.cantidad_compresores),
    ]


def _datos_temperaturas(detalle: SmuDetallePreventivoAa | None) -> list:
    if detalle is None:
        return []
    return [
        base.numero(detalle.temp_cuarto),
        base.numero(detalle.temp_entrada),
        base.numero(detalle.temp_salida),
        base.numero(detalle.temp_display),
        texto_o_vacio(detalle.mp_termostato),
        texto_o_vacio(detalle.ajuste_termostato),
        texto_o_vacio(detalle.mp_termostato_post),
    ]


def _triple(a, b, c) -> str:
    """``VL1 / VL2 / VL3`` omitiendo los que vienen en 0 (sin diligenciar)."""
    partes = [base.numero(valor) for valor in (a, b, c)]
    return " / ".join(p for p in partes if p) or ""


def _fila_etiquetada(ws, fila: int, etiqueta: str, grupos, valores) -> int:
    """Etiqueta rosa en A + celdas amarillas agrupadas (mini tablas)."""
    celda = ws.cell(row=fila, column=1)
    celda.value = etiqueta
    celda.fill = FILL_ROSA
    celda.font = FONT_DATA_BOLD
    celda.border = BORDER_TABLA
    return base.datos_grupos(ws, fila, grupos, valores)


# ── Secciones ────────────────────────────────────────────────────────────────

def _seccion_informacion_general(ws, fila: int, actividad: SmuActividad) -> int:
    pares = pares_con_datos(_pares_informacion_general(actividad))
    if not pares:
        return fila
    fila_banner(ws, fila, "1. INFORMACIÓN GENERAL", FILL_BURGUNDY, FONT_HEADER_WHITE)
    fila += 1
    base.fila_espaciadora(ws, fila); fila += 1
    for par in pares:
        par_de_campos(ws, fila, par[0], par[1])
        fila += 1
    base.fila_espaciadora(ws, fila); fila += 1
    return fila


def _seccion_informacion_aa(
    ws, fila: int, actividad: SmuActividad, detalle
) -> int:
    compresores = list(actividad.aa_compresores.all())
    condensadoras = list(actividad.aa_condensadoras.all())
    manejadoras = list(actividad.aa_manejadoras.all())
    if detalle is None and not (compresores or condensadoras or manejadoras):
        return fila

    fila_banner(ws, fila, "1.- INFORMACIÓN AIRE ACONDICIONADO",
                FILL_BURGUNDY, FONT_HEADER_WHITE)
    fila += 1
    base.fila_espaciadora(ws, fila); fila += 1

    # Datos generales del equipo (una fila por detalle registrado) + temperaturas.
    if detalle is not None:
        fila = sub_titulo(ws, fila, "DATOS GENERALES AIRES ACONDICIONADOS")
        base.fila_espaciadora(ws, fila); fila += 1

        fila_encabezados_tabla(ws, fila, ENCABEZADOS_EQUIPO, col_ini=2)
        ws.cell(row=fila, column=1).fill = FILL_ROSA
        fila += 1
        fila = fila_con_etiqueta(ws, fila, "AA. 1", len(ENCABEZADOS_EQUIPO))
        for columna, valor in enumerate(_datos_equipo(detalle), start=2):
            ws.cell(row=fila - 1, column=columna).value = valor or None

        # El ID de activo no tiene columna en la plantilla: va en su propia fila.
        id_activo = texto_o_vacio(detalle.id_activo).strip()
        if id_activo:
            par_de_campos(ws, fila, ("ID ACTIVO:", id_activo), ("", ""))
            fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

        encabezados_grupos(ws, fila, GRUPOS_7, ENCABEZADOS_TEMPERATURAS,
                           fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE)
        fila += 1
        fila = base.datos_grupos(ws, fila, GRUPOS_7, _datos_temperaturas(detalle))
        base.fila_espaciadora(ws, fila); fila += 1

    if compresores:
        fila = sub_titulo(ws, fila, "UNIDAD CONDENSADORA")
        base.fila_espaciadora(ws, fila); fila += 1

        fila_encabezados_tabla(ws, fila, ENCABEZADOS_COMPRESOR, col_ini=2)
        ws.cell(row=fila, column=1).fill = FILL_ROSA
        fila += 1
        for indice, compresor in enumerate(compresores, start=1):
            valores = [
                texto_o_vacio(compresor.marca),
                texto_o_vacio(compresor.serial),
                texto_o_vacio(compresor.refrigerante),
                texto_o_vacio(compresor.modelo),
                texto_o_vacio(compresor.aislamiento_electrico),
                base.numero(compresor.presion_succion),
                base.numero(compresor.presion_descarga),
                texto_o_vacio(compresor.nivel_aceite),
                _triple(compresor.vl1, compresor.vl2, compresor.vl3),
                _triple(compresor.amp_l1, compresor.amp_l2, compresor.amp_l3),
            ]
            fila = fila_con_etiqueta(ws, fila, f"COMPRESOR {indice}",
                                     len(ENCABEZADOS_COMPRESOR))
            for columna, valor in enumerate(valores, start=2):
                ws.cell(row=fila - 1, column=columna).value = valor or None
        base.fila_espaciadora(ws, fila); fila += 1

    if condensadoras:
        encabezados_grupos(ws, fila, GRUPOS_CONDENSADORA, ENCABEZADOS_CONDENSADORA,
                           fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE)
        ws.cell(row=fila, column=1).fill = FILL_GRIS_MEDIO
        ws.cell(row=fila, column=1).border = BORDER_NEGRO
        fila += 1
        for indice, condensadora in enumerate(condensadoras, start=1):
            valores = [
                texto_o_vacio(condensadora.marca),
                texto_o_vacio(condensadora.modelo),
                texto_o_vacio(condensadora.serial),
                base.numero(condensadora.temperatura_entrada),
                base.numero(condensadora.temperatura_salida),
                texto_o_vacio(condensadora.diametro_eje),
                texto_o_vacio(condensadora.diametro_aspas),
            ]
            fila = _fila_etiquetada(ws, fila, f"CONDENSADORA {indice}",
                                    GRUPOS_CONDENSADORA, valores)
        base.fila_espaciadora(ws, fila); fila += 1

    if manejadoras:
        fila = sub_titulo(ws, fila, "UNIDAD MANEJADORA")
        base.fila_espaciadora(ws, fila); fila += 1

        fila_encabezados_tabla(ws, fila, ENCABEZADOS_MANEJADORA, col_ini=2)
        ws.cell(row=fila, column=1).fill = FILL_ROSA
        fila += 1
        for indice, manejadora in enumerate(manejadoras, start=1):
            valores = [
                texto_o_vacio(manejadora.marca),
                texto_o_vacio(manejadora.modelo),
                texto_o_vacio(manejadora.tipo),
                texto_o_vacio(manejadora.tipo_filtro),
                texto_o_vacio(manejadora.tipo_correa),
                texto_o_vacio(manejadora.marca_motor),
                texto_o_vacio(manejadora.alimentacion_ac),
                base.numero(manejadora.voltaje),
                base.numero(manejadora.corriente),
                base.numero(manejadora.aislamiento),
                texto_o_vacio(manejadora.serial_motor),
            ]
            fila = fila_con_etiqueta(ws, fila, f"MANEJADORA {indice}",
                                     len(ENCABEZADOS_MANEJADORA))
            for columna, valor in enumerate(valores, start=2):
                ws.cell(row=fila - 1, column=columna).value = valor or None
        base.fila_espaciadora(ws, fila); fila += 1

    return fila


def _seccion_checklist(ws, fila: int, actividad: SmuActividad) -> int:
    items = list(actividad.checklist_aa.order_by("item_numero"))
    if not items:
        return fila

    fila_banner(ws, fila, "2.- LISTA DE CHEQUEO GENERAL",
                FILL_BURGUNDY, FONT_HEADER_WHITE)
    fila += 1
    base.fila_espaciadora(ws, fila); fila += 1

    fusionar(ws, fila, 1, fila, 11, valor="PARAMETRO DE EVALUACIÓN",
             fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
    fusionar(ws, fila, 12, fila, 12, valor="EVAL",
             fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
    fila += 1

    for item in items:
        fusionar(ws, fila, 1, fila, 11, valor=parametro_de(item.item_numero),
                 font=FONT_DATA_REGULAR, border=BORDER_TABLA,
                 align=base.ALIGN_LEFT)
        fusionar(ws, fila, 12, fila, 12, valor=texto_o_vacio(item.estado) or None,
                 fill=base.FILL_YELLOW, font=FONT_DATA_REGULAR,
                 border=BORDER_TABLA, align=base.ALIGN_CENTER)
        fila += 1
    base.fila_espaciadora(ws, fila); fila += 1
    return fila


def _seccion_acciones(ws, fila: int, actividad: SmuActividad) -> int:
    acciones = list(actividad.aa_acciones.all())
    if not acciones:
        return fila

    fila_banner(ws, fila, "3.- ACCIONES REALIZADAS", FILL_BURGUNDY, FONT_HEADER_WHITE)
    fila += 1
    base.fila_espaciadora(ws, fila); fila += 1

    fila_encabezados_tabla(ws, fila, [t for _, t in ACCIONES_AA], col_ini=2)
    ws.cell(row=fila, column=1).fill = FILL_ROSA
    fila += 1

    # Una fila por cada unidad AA (aa_acciones.orden), como en la plantilla.
    por_orden: dict[int, dict[str, str]] = {}
    for accion in acciones:
        por_orden.setdefault(accion.orden, {})[accion.clave] = accion.valor

    for orden in sorted(por_orden):
        valores = [
            VALOR_LEGIBLE.get(por_orden[orden].get(clave, ""), "")
            for clave, _ in ACCIONES_AA
        ]
        fila = fila_con_etiqueta(ws, fila, f"AA {orden}", len(ACCIONES_AA))
        for columna, valor in enumerate(valores, start=2):
            ws.cell(row=fila - 1, column=columna).value = valor or None
    base.fila_espaciadora(ws, fila); fila += 1
    return fila


def _seccion_observaciones(ws, fila: int, detalle) -> int:
    observaciones = texto_o_vacio(detalle.acciones_observaciones if detalle else "")
    plan_mejora = texto_o_vacio(detalle.plan_mejora if detalle else "")
    if not (observaciones.strip() or plan_mejora.strip()):
        return fila

    fila_banner(ws, fila, "4.- OBSERVACIONES Y PLAN DE MEJORA",
                FILL_BURGUNDY, FONT_HEADER_WHITE)
    fila += 1
    base.fila_espaciadora(ws, fila); fila += 1

    if observaciones.strip():
        fusionar(ws, fila, 1, fila, 12, valor="OBSERVACIONES DE LAS ACCIONES",
                 fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
        fila += 1
        fila = bloque_texto(ws, fila, observaciones, 1, 12)
        base.fila_espaciadora(ws, fila); fila += 1

    if plan_mejora.strip():
        fusionar(ws, fila, 1, fila, 12, valor="PLAN DE MEJORA",
                 fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
        fila += 1
        fila = bloque_texto(ws, fila, plan_mejora, 1, 12)
        base.fila_espaciadora(ws, fila); fila += 1
    return fila


# ── Generador ────────────────────────────────────────────────────────────────

def construir_libro(actividad: SmuActividad) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    base.aplicar_anchos(ws)

    detalle = _detalle(actividad)
    evidencias = list(actividad.evidencias.all())

    fila = 1

    # ── Encabezado del documento ────────────────────────────────────────────
    base.fila_espaciadora(ws, fila); fila += 1
    fila_banner(ws, fila, "DIRECCIÓN DE O&M RED DE ACCESO",
                FILL_BURGUNDY, FONT_HEADER_WHITE); fila += 1
    fila_banner(ws, fila, "PLAN DE MANTENIMIENTO PREVENTIVO INTEGRAL RED MOVIL",
                FILL_BURGUNDY, FONT_HEADER_WHITE); fila += 1
    fila_banner(ws, fila, "MANTENIMIENTO PREVENTIVO AIRES ACONDICIONADOS",
                FILL_BURGUNDY, FONT_HEADER_WHITE); fila += 1
    base.fila_espaciadora(ws, fila); fila += 1

    fila = _seccion_informacion_general(ws, fila, actividad)
    fila = _seccion_informacion_aa(ws, fila, actividad, detalle)
    fila = _seccion_checklist(ws, fila, actividad)
    fila = _seccion_acciones(ws, fila, actividad)

    # Numeración dinámica: la plantilla cierra en "3.-"; los bloques que no
    # existen en el formato original se agregan después.
    numero = 4
    con_observaciones = bool(
        detalle and (texto_o_vacio(detalle.acciones_observaciones).strip()
                     or texto_o_vacio(detalle.plan_mejora).strip())
    )
    if con_observaciones:
        fila = _seccion_observaciones(ws, fila, detalle)
        numero += 1

    if evidencias:
        fila_banner(ws, fila, f"{numero}.- ANEXOS FOTOGRÁFICOS",
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1
        fila = bloque_fotos(ws, fila, evidencias, titulo=base.rotulo_evidencia)
        numero += 1

    # Certificados, firmas y fotos de técnicos: imágenes que no son evidencias
    # y que ningún formato pintaba (sin foto no se pinta la sección).
    fila = base.bloque_imagenes_adjuntas(
        ws, fila, actividad,
        titulo=f"{numero}.- CERTIFICADOS, TÉCNICOS Y FIRMAS")

    estilizar_fila(ws, 1, border=BORDER_NEGRO)
    return wb


def generar_excel_preventivo_aa(actividad: SmuActividad) -> Workbook:
    """Devuelve el libro Excel del preventivo de aire acondicionado."""
    return construir_libro(actividad)
