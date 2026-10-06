"""Exporte de los correctivos CAPEX (tipologías 1, 3 y 5) a Excel.

Port dinámico de ``plantillaSmu.py`` (hoja "Tipologia SPT"): banners
borgoña, metadatos de la tipología 1, tabla de inventario con columna de
foto, registro fotográfico, transportes, resultados SPT, panorámica/planos,
otras actividades (rosa magenta) y recomendaciones.

Las secciones se construyen según los datos reales:

- una fila de inventario por material (no el catálogo fijo de la plantilla),
- una fila de mano de obra por actividad, con sus fotos antes/después,
- una fila de transporte con sus fotos antes/durante/después,
- la tabla Wenner de resistividad en lugar de campos que el formulario no pide,
- las tipologías 3 y 5 reutilizan el mismo estilo con sus secciones aplicables
  (no traen plantilla propia ni datos SPT, así que esas filas no se pintan),
- ninguna foto se pierde: lo que no tiene sección dedicada cae en el
  "registro fotográfico".
"""
from __future__ import annotations

import re

from openpyxl import Workbook

from apps.smu.models import SmuActividad, SmuDetalleCapex

from . import base
from .base import (
    ALIGN_LEFT,
    BORDER_NEGRO,
    BORDER_TABLA,
    FILL_BURGUNDY,
    FILL_GRIS_MEDIO,
    FILL_LIGHT_GRAY,
    FILL_PINK_MAGENTA,
    FILL_WHITE,
    FONT_DATA_BOLD,
    FONT_DATA_REGULAR,
    FONT_HEADER_WHITE,
    Numerador,
    bloque_fotos,
    bloque_texto,
    datos_grupos,
    encabezados_grupos,
    estilizar_fila,
    fila_banner,
    fusionar,
    par_de_campos,
    pares_con_datos,
    sub_titulo,
    texto_o_vacio,
)
from .imagenes import insertar_foto_en_celda

HOJAS = {
    "tipologia1": "Tipologia SPT",
    "tipologia3": "Tipologia 3",
    "tipologia5": "Tipologia 5",
}

TITULOS_TIPOLOGIA = {
    "tipologia1": ("MANTENIMIENTO CORRECTIVO DE TIPOLOGIA - 1 - "
                   "SISTEMA DE PUESTA A TIERRA"),
    "tipologia3": "MANTENIMIENTO CORRECTIVO DE TIPOLOGIA - 3 - CLIMATIZACIÓN",
    "tipologia5": ("MANTENIMIENTO CORRECTIVO DE TIPOLOGIA - 5 - "
                   "SUBESTACIONES ELÉCTRICAS"),
}

# Metadatos normativos: solo la plantilla de la tipología 1 los trae.
DEFINICIONES_T1 = [
    ("Definición",
     "Es un mecanismo de seguridad que forma parte de las instalaciones "
     "eléctricas y que consiste en conducir eventuales desvíos de la corriente "
     "hacia la tierra, impidiendo que el usuario y equipos electromecánicos "
     " entren en contacto con la electricidad."),
    ("Caso de uso", "Reposición por Hurto Renovación por alto deterioro"),
    ("Alcance",
     "Solo aplica para instalación de sistemas completos de puesta a tierra."),
]

# ── Tablas (grupos de columnas sobre A..L) ───────────────────────────────────
GRUPOS_INVENTARIO = [
    (1, 2), (3, 3), (4, 5), (6, 7), (8, 8), (9, 9), (10, 10), (11, 12),
]
ENCABEZADOS_INVENTARIO = [
    "TEXTO SAP", "CODIGO SAP", "ALCANCE", "COMENTARIOS",
    "CANTIDAD ESTANDAR", "CANTIDAD REAL", "UNIDAD DE MEDIDA", "FOTO",
]

GRUPOS_MANO_OBRA = [(1, 4), (5, 5), (6, 6), (7, 8), (9, 10), (11, 12)]
ENCABEZADOS_MANO_OBRA = [
    "DESCRIPCIÓN DE ACTIVIDAD EXTRA", "CANT", "UND",
    "FOTO ANTES", "FOTO DURANTE", "FOTO DESPUÉS",
]

GRUPOS_TRANSPORTE = [(1, 3), (4, 4), (5, 7), (8, 10), (11, 12)]
ENCABEZADOS_TRANSPORTE = [
    "DESCRIPCION Y CARACTERISTICAS DEL TRANSPORTE", "CODIGO SAP",
    "FOTO ANTES", "FOTO DURANTE", "FOTO DESPUES",
]

GRUPOS_SPT = [(1, 4), (5, 8), (9, 12)]
ENCABEZADOS_SPT = ["DISTANCIA", "MEDIDA (OHMIOS)", "RESISTIVIDAD"]

# Altura de una fila con foto (la caja máxima es 140 px ≈ 105 pt).
ALTO_CON_FOTO = 105
ALTO_FILA_TABLA = 30

# Secciones que sus fotos tienen destino propio (no van al registro general).
SECCIONES_CON_DESTINO = {
    "inventario", "mano_obra", "transporte", "spt", "panoramica",
    "matricula", "justificacion",
}


# ── Utilidades de fotos ──────────────────────────────────────────────────────

def _evidencia(evidencias, seccion: str, campo: str):
    for e in evidencias:
        if e.seccion == seccion and e.campo_origen == campo:
            return e
    return None


def _foto_inventario(evidencias, material, posicion: int):
    for e in evidencias:
        if e.seccion == "inventario" and e.material_id == material.id:
            return e
    for e in evidencias:
        if e.seccion != "inventario":
            continue
        coincidencia = re.match(r"^insumo_(\d+)_foto$", e.campo_origen)
        if coincidencia and int(coincidencia.group(1)) == posicion:
            return e
    return None


def _foto_mano_obra(evidencias, posicion: int, parte: str):
    """`parte` ∈ ``antes`` | ``despues``."""
    for e in evidencias:
        if e.seccion != "mano_obra":
            continue
        coincidencia = re.match(rf"^mo_(\d+)_foto_{parte}$", e.campo_origen)
        if coincidencia and int(coincidencia.group(1)) == posicion:
            return e
    return None


def _foto_transporte(evidencias, posicion: int, parte: str):
    """`parte` ∈ ``antes`` | ``durante`` | ``despues``.

    La tipología 1 guarda una sola fila de transporte sin índice
    (``transporte_foto_antes``); las demás usan ``transporte_N_foto_...``.
    """
    for e in evidencias:
        if e.seccion != "transporte":
            continue
        if e.campo_origen == f"transporte_foto_{parte}" and posicion == 1:
            return e
        coincidencia = re.match(rf"^transporte_(\d+)_foto_{parte}$", e.campo_origen)
        if coincidencia and int(coincidencia.group(1)) == posicion:
            return e
    return None


def _caja_foto(ws, fila: int, titulo: str, evidencia) -> int:
    """Caja de foto a todo el ancho (panorámica, telurómetro, matrícula)."""
    fusionar(ws, fila, 1, fila, 12, valor=titulo, fill=FILL_BURGUNDY,
             font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
    fila += 1
    ancla = fusionar(ws, fila, 1, fila + 5, 12, fill=FILL_LIGHT_GRAY,
                     font=FONT_DATA_REGULAR, border=BORDER_TABLA)
    for r in range(fila, fila + 6):
        ws.row_dimensions[r].height = 25
    if evidencia is not None:
        insertar_foto_en_celda(ws, evidencia.ruta_archivo, ancla)
    fila += 6
    base.fila_espaciadora(ws, fila)
    return fila + 1


def _bloque_fotos_completo(ws, fila: int, evidencias) -> int:
    """Media caja por foto (dos por fila), como el registro fotográfico."""
    if len(evidencias) == 1:
        return _caja_foto(ws, fila, base.rotulo_evidencia(evidencias[0]),
                          evidencias[0])
    return bloque_fotos(ws, fila, evidencias, titulo=base.rotulo_evidencia)


# ── Datos de la actividad ────────────────────────────────────────────────────

def _detalle(actividad: SmuActividad) -> SmuDetalleCapex | None:
    try:
        return actividad.detalle_capex
    except SmuDetalleCapex.DoesNotExist:
        return None


def _pares_informacion_general(actividad: SmuActividad, detalle):
    izquierda = [
        ("NOMBRE ESTACION BASE:", texto_o_vacio(actividad.nombre_estacion)),
        ("TIPO ESTACION BASE:", texto_o_vacio(actividad.tipo_estacion)),
        ("NOMBRE SITE OWNER:", texto_o_vacio(actividad.site_owner)),
        ("NOMBRE DE QUIEN EJECUTA ACTIVIDAD:", base.responsable(actividad)),
        ("COORDINADOR ALIADO:", texto_o_vacio(actividad.coordinador_aliado)),
        ("NOMBRE DE LA MATRÍCULA:",
         texto_o_vacio(detalle.matricula_nombre if detalle else "")),
        ("FECHA DE LA MATRÍCULA:",
         base.formato_fecha(detalle.matricula_fecha) if detalle else ""),
    ]
    derecha = [
        ("OT:", texto_o_vacio(actividad.codigo_ot)),
        ("FECHA INICIO ACTIVIDAD:",
         base.formato_fecha(actividad.fecha_inicio, con_hora=True)),
        ("FECHA FINAL ACTIVIDAD:",
         base.formato_fecha(actividad.fecha_fin, con_hora=True)),
        ("EMPRESA / ALIADO:", texto_o_vacio(actividad.empresa)),
        ("ESTADO:", actividad.get_estado_display()),
        ("No. DE MATRÍCULA:", texto_o_vacio(detalle.matricula_numero if detalle else "")),
        ("", ""),
    ]
    return list(zip(izquierda, derecha))


# ── Generador ────────────────────────────────────────────────────────────────

def construir_libro(actividad: SmuActividad) -> Workbook:
    detalle = _detalle(actividad)
    tipologia = (detalle.tipologia if detalle else actividad.tipo_formulario) \
        or "tipologia1"

    wb = Workbook()
    ws = wb.active
    ws.title = HOJAS.get(tipologia, "Tipologia")
    base.aplicar_anchos(ws)

    materiales = list(actividad.materiales.all())
    inventario = [m for m in materiales if m.tipo_registro == "inventario_capex"]
    mano_obra = [m for m in materiales if m.tipo_registro == "mano_obra_capex"]
    transportes = list(actividad.transportes.all())
    spt_filas = list(actividad.spt_filas.all())
    evidencias = list(actividad.evidencias.all())
    usadas: set[int] = set()

    numeros = Numerador()
    fila = 1

    # ── Encabezado del documento ────────────────────────────────────────────
    base.fila_espaciadora(ws, fila); fila += 1
    fila_banner(ws, fila, "DIRECCIÓN DE O&M RED DE ACCESO",
                FILL_BURGUNDY, FONT_HEADER_WHITE); fila += 1
    fila_banner(
        ws, fila,
        "PLAN DE MANTENIMIENTO CORRECTIVO CAPITALIZABLE INTEGRAL RED MOVIL",
        FILL_BURGUNDY, FONT_HEADER_WHITE,
    ); fila += 1
    fila_banner(ws, fila, TITULOS_TIPOLOGIA.get(tipologia, tipologia),
                FILL_BURGUNDY, FONT_HEADER_WHITE); fila += 1
    base.fila_espaciadora(ws, fila); fila += 1

    # ── Metadatos normativos (solo tipología 1) ─────────────────────────────
    if tipologia == "tipologia1":
        for etiqueta, texto in DEFINICIONES_T1:
            celda = ws.cell(row=fila, column=1, value=etiqueta)
            celda.font = FONT_DATA_BOLD
            celda.fill = FILL_LIGHT_GRAY
            celda.alignment = ALIGN_LEFT
            celda.border = base.BORDER_THIN
            fusionar(ws, fila, 2, fila, 12, valor=texto, fill=FILL_LIGHT_GRAY,
                     font=FONT_DATA_REGULAR, border=base.BORDER_THIN,
                     align=ALIGN_LEFT)
            fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

    # ── 1. Información general ──────────────────────────────────────────────
    pares_general = pares_con_datos(_pares_informacion_general(actividad, detalle))
    if pares_general:
        fila_banner(ws, fila, numeros.titulo("INFORMACION GENERAL"),
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1
        for par in pares_general:
            par_de_campos(ws, fila, par[0], par[1])
            fila += 1

        soporte_matricula = _evidencia(evidencias, "matricula",
                                       "foto_soporte_matricula_preview")
        if soporte_matricula is not None:
            usadas.add(soporte_matricula.id)
            fila = _caja_foto(ws, fila, "SOPORTE DE MATRÍCULA PROFESIONAL",
                              soporte_matricula)
        else:
            base.fila_espaciadora(ws, fila); fila += 1

    # ── 2. Inventario de suministros (1 fila por material) ──────────────────
    if inventario:
        fila_banner(ws, fila, numeros.titulo("INVENTARIO DE SUMINISTROS"),
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

        fila = encabezados_grupos(ws, fila, GRUPOS_INVENTARIO,
                                  ENCABEZADOS_INVENTARIO)
        for posicion, material in enumerate(inventario, start=1):
            foto = _foto_inventario(evidencias, material, posicion)
            if foto is not None:
                usadas.add(foto.id)
            valores = [
                texto_o_vacio(material.texto_sap),
                texto_o_vacio(material.codigo_sap),
                texto_o_vacio(material.alcance),
                texto_o_vacio(material.comentarios),
                base.numero(material.cantidad_estandar),
                base.numero(material.cantidad_real),
                texto_o_vacio(material.unidad_medida),
                None,
            ]
            fila = datos_grupos(ws, fila, GRUPOS_INVENTARIO, valores,
                                fill=FILL_WHITE, align=ALIGN_LEFT)
            ws.row_dimensions[fila - 1].height = (
                ALTO_CON_FOTO if foto is not None else ALTO_FILA_TABLA
            )
            if foto is not None:
                ancla = ws.cell(row=fila - 1, column=11).coordinate
                insertar_foto_en_celda(ws, foto.ruta_archivo, ancla)
        base.fila_espaciadora(ws, fila); fila += 1

    # ── 3. Registro fotográfico (fotos sin sección propia) ──────────────────
    pendientes = [e for e in evidencias if e.id not in usadas
                  and e.seccion not in SECCIONES_CON_DESTINO]
    secciones = []
    for e in pendientes:
        if e.seccion not in secciones:
            secciones.append(e.seccion)
    if secciones:
        fila_banner(
            ws, fila,
            numeros.titulo("REGISTRO FOTOGRAFICO DE LA ACTIVIDAD "
                           "(EVIDENCIAS DE EJECUCION)"),
            FILL_BURGUNDY, FONT_HEADER_WHITE,
        )
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1
        for seccion in secciones:
            fotos = [e for e in pendientes if e.seccion == seccion]
            for foto in fotos:
                usadas.add(foto.id)
            fila = sub_titulo(ws, fila, seccion.upper())
            fila = _bloque_fotos_completo(ws, fila, fotos)

    # ── 4. Mano de obra / otras actividades aprobadas ───────────────────────
    otras_actividades = [
        texto_o_vacio(detalle.otras_actividades_1 if detalle else ""),
        texto_o_vacio(detalle.otras_actividades_2 if detalle else ""),
        texto_o_vacio(detalle.otras_actividades_3 if detalle else ""),
    ]
    justificacion = texto_o_vacio(
        detalle.texto_justificacion_otras_actividades if detalle else ""
    )
    fotos_justificacion = [e for e in evidencias
                           if e.seccion == "justificacion"
                           and e.id not in usadas]
    hay_otras = bool(mano_obra or any(t.strip() for t in otras_actividades)
                     or justificacion.strip() or fotos_justificacion)
    if hay_otras:
        fila_banner(
            ws, fila,
            numeros.titulo("RESULTADOS DE OTRAS ACTIVIDADES APROBADAS "
                           "POR FUERA DEL ESTANDAR."),
            FILL_PINK_MAGENTA, FONT_HEADER_WHITE,
        )
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

        if mano_obra:
            fila = encabezados_grupos(ws, fila, GRUPOS_MANO_OBRA,
                                      ENCABEZADOS_MANO_OBRA,
                                      fill=FILL_PINK_MAGENTA)
            for posicion, material in enumerate(mano_obra, start=1):
                antes = _foto_mano_obra(evidencias, posicion, "antes")
                despues = _foto_mano_obra(evidencias, posicion, "despues")
                for foto in (antes, despues):
                    if foto is not None:
                        usadas.add(foto.id)
                descripcion = base.descripcion_material(material) \
                    or texto_o_vacio(material.alcance)
                cant = base.numero(material.cantidad_real)
                valores = [
                    descripcion,
                    cant,
                    texto_o_vacio(material.unidad_medida),
                    None, None, None,
                ]
                fila = datos_grupos(ws, fila, GRUPOS_MANO_OBRA, valores,
                                    fill=FILL_WHITE, align=ALIGN_LEFT)
                ws.row_dimensions[fila - 1].height = (
                    ALTO_CON_FOTO if (antes or despues) else ALTO_FILA_TABLA
                )
                if antes is not None:
                    insertar_foto_en_celda(
                        ws, antes.ruta_archivo,
                        ws.cell(row=fila - 1, column=7).coordinate,
                    )
                if despues is not None:
                    insertar_foto_en_celda(
                        ws, despues.ruta_archivo,
                        ws.cell(row=fila - 1, column=11).coordinate,
                    )

        textos_extra = [t for t in otras_actividades if t.strip()]
        if textos_extra:
            numerados = "\n".join(
                f"{i}. {t}" for i, t in enumerate(textos_extra, start=1)
            )
            fusionar(ws, fila, 1, fila, 12, valor="OTRAS ACTIVIDADES",
                     fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE,
                     border=BORDER_NEGRO)
            fila += 1
            fila = bloque_texto(ws, fila, numerados, 1, 12)

        if justificacion.strip():
            fusionar(ws, fila, 1, fila, 12, valor="JUSTIFICACIÓN",
                     fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE,
                     border=BORDER_NEGRO)
            fila += 1
            fila = bloque_texto(ws, fila, justificacion, 1, 12)

        if fotos_justificacion:
            for foto in fotos_justificacion:
                usadas.add(foto.id)
            base.fila_espaciadora(ws, fila); fila += 1
            fila = _bloque_fotos_completo(ws, fila, fotos_justificacion)
        base.fila_espaciadora(ws, fila); fila += 1

    # ── 5. Transportes ──────────────────────────────────────────────────────
    if transportes:
        fila_banner(ws, fila, numeros.titulo("TRANSPORTES"),
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

        fila = encabezados_grupos(ws, fila, GRUPOS_TRANSPORTE,
                                  ENCABEZADOS_TRANSPORTE)
        for posicion, transporte in enumerate(transportes, start=1):
            antes = _foto_transporte(evidencias, posicion, "antes")
            durante = _foto_transporte(evidencias, posicion, "durante")
            despues = _foto_transporte(evidencias, posicion, "despues")
            for foto in (antes, durante, despues):
                if foto is not None:
                    usadas.add(foto.id)

            partes = [
                texto_o_vacio(transporte.descripcion),
                texto_o_vacio(transporte.tipo_transporte),
            ]
            distancia = transporte.distancia_km
            if distancia not in (None, 0):
                partes.append(f"{float(distancia):g} km")
            if texto_o_vacio(transporte.tiempo_traslado):
                partes.append(texto_o_vacio(transporte.tiempo_traslado))
            if texto_o_vacio(transporte.observacion):
                partes.append(f"Obs: {texto_o_vacio(transporte.observacion)}")
            descripcion = " | ".join(p for p in partes if p)

            valores = [
                descripcion,
                texto_o_vacio(transporte.codigo_sap),
                None, None, None,
            ]
            fila = datos_grupos(ws, fila, GRUPOS_TRANSPORTE, valores,
                                fill=FILL_WHITE, align=ALIGN_LEFT)
            ws.row_dimensions[fila - 1].height = (
                ALTO_CON_FOTO if (antes or durante or despues)
                else ALTO_FILA_TABLA
            )
            for columna, foto in ((5, antes), (8, durante), (11, despues)):
                if foto is not None:
                    insertar_foto_en_celda(
                        ws, foto.ruta_archivo,
                        ws.cell(row=fila - 1, column=columna).coordinate,
                    )
        base.fila_espaciadora(ws, fila); fila += 1

    # ── 6. Resultados de reforma SPT (solo tipología 1, con datos) ──────────
    spt_foto = _evidencia(evidencias, "spt", "spt_naturaleza_terreno_foto")
    spt_naturaleza = texto_o_vacio(detalle.spt_naturaleza if detalle else "")
    spt_recomendacion = texto_o_vacio(detalle.spt_recomendacion if detalle else "")
    if (tipologia == "tipologia1"
            and (spt_filas or spt_naturaleza.strip() or spt_recomendacion.strip()
                 or spt_foto is not None)):
        fila_banner(ws, fila, numeros.titulo("RESULTADOS DE REFORMA SPT"),
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

        if spt_filas:
            fila = encabezados_grupos(ws, fila, GRUPOS_SPT, ENCABEZADOS_SPT)
            for fila_spt in spt_filas:
                valores = [
                    texto_o_vacio(fila_spt.distancia),
                    texto_o_vacio(fila_spt.medida_ohmio),
                    texto_o_vacio(fila_spt.resistividad),
                ]
                fila = datos_grupos(ws, fila, GRUPOS_SPT, valores,
                                    fill=FILL_WHITE, align=ALIGN_LEFT)
                ws.row_dimensions[fila - 1].height = ALTO_FILA_TABLA
            base.fila_espaciadora(ws, fila); fila += 1

        if spt_naturaleza.strip():
            fusionar(ws, fila, 1, fila, 12,
                     valor="NATURALEZA DEL TERRENO (TELURÓMETRO)",
                     fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE,
                     border=BORDER_NEGRO)
            fila += 1
            fila = bloque_texto(ws, fila, spt_naturaleza, 1, 12)
            base.fila_espaciadora(ws, fila); fila += 1

        if spt_foto is not None:
            usadas.add(spt_foto.id)
            fila = _caja_foto(ws, fila, "FOTO NATURALEZA DEL TERRENO",
                              spt_foto)

        if spt_recomendacion.strip():
            fusionar(ws, fila, 1, fila, 12, valor="RECOMENDACIÓN AL ALIADO",
                     fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE,
                     border=BORDER_NEGRO)
            fila += 1
            fila = bloque_texto(ws, fila, spt_recomendacion, 1, 12)
            base.fila_espaciadora(ws, fila); fila += 1

    # ── 7. Panorámica y planos ──────────────────────────────────────────────
    panoramica = _evidencia(evidencias, "panoramica", "foto_panoramica_preview")
    plano = _evidencia(evidencias, "panoramica", "foto_plano_preview")
    recomendacion_panoramica = texto_o_vacio(
        detalle.panoramica_recomendacion if detalle else ""
    )
    if panoramica is not None or plano is not None or recomendacion_panoramica.strip():
        fila_banner(
            ws, fila,
            numeros.titulo("PANORAMICA DE LA TORRE Y PLANOS DE TERRENO "
                           "EN MANO ALZADA"),
            FILL_BURGUNDY, FONT_HEADER_WHITE,
        )
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1

        cajas = [e for e in (panoramica, plano) if e is not None]
        for foto in cajas:
            usadas.add(foto.id)
        if cajas:
            if len(cajas) == 2:
                fusionar(ws, fila, 1, fila, 6,
                         valor="PANORÁMICA DE LA TORRE / ESTACIÓN BASE",
                         fill=FILL_BURGUNDY, font=FONT_HEADER_WHITE,
                         border=BORDER_NEGRO)
                fusionar(ws, fila, 7, fila, 12,
                         valor="PLANO DE TERRENO A MANO ALZADA (EB)",
                         fill=FILL_BURGUNDY, font=FONT_HEADER_WHITE,
                         border=BORDER_NEGRO)
                fila += 1
                anclas = [
                    fusionar(ws, fila, lado_ini, fila + 5, lado_fin,
                             fill=FILL_LIGHT_GRAY, font=FONT_DATA_REGULAR,
                             border=BORDER_TABLA)
                    for lado_ini, lado_fin in ((1, 6), (7, 12))
                ]
                for r in range(fila, fila + 6):
                    ws.row_dimensions[r].height = 25
                for ancla, foto in zip(anclas, cajas):
                    insertar_foto_en_celda(ws, foto.ruta_archivo, ancla)
                fila += 6
            else:
                rotulo = ("PLANO DE TERRENO A MANO ALZADA (EB)"
                          if "plano" in cajas[0].campo_origen
                          else "PANORÁMICA DE LA TORRE / ESTACIÓN BASE")
                fila = _caja_foto(ws, fila, rotulo, cajas[0])

        if recomendacion_panoramica.strip():
            fusionar(ws, fila, 1, fila, 12,
                     valor="RECOMENDACIÓN DEL ALIADO (PANORÁMICA)",
                     fill=FILL_GRIS_MEDIO, font=FONT_HEADER_WHITE,
                     border=BORDER_NEGRO)
            fila += 1
            fila = bloque_texto(ws, fila, recomendacion_panoramica, 1, 12)
            base.fila_espaciadora(ws, fila); fila += 1

    # ── 8. Descripción general y recomendaciones ────────────────────────────
    bloques = [
        ("DESCRIPCIÓN GENERAL DE LAS ACTIVIDADES",
         texto_o_vacio(detalle.descripcion_general_actividades if detalle else "")),
        ("RECOMENDACIONES DEL ALIADO",
         texto_o_vacio(detalle.recomendaciones_aliado if detalle else "")),
        ("RECOMENDACIÓN FINAL",
         texto_o_vacio(detalle.recomendacion_final if detalle else "")),
        ("OBSERVACIONES", texto_o_vacio(actividad.observaciones)),
        ("HALLAZGOS", texto_o_vacio(actividad.hallazgos)),
        ("RECOMENDACIONES", texto_o_vacio(actividad.recomendaciones)),
    ]
    bloques = [(rotulo, texto) for rotulo, texto in bloques if texto.strip()]
    if bloques:
        fila_banner(ws, fila, numeros.titulo("RECOMENDACIONES DE LA ACTIVIDAD"),
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1
        for rotulo, texto in bloques:
            fusionar(ws, fila, 1, fila, 12, valor=rotulo, fill=FILL_GRIS_MEDIO,
                     font=FONT_HEADER_WHITE, border=BORDER_NEGRO)
            fila += 1
            fila = bloque_texto(ws, fila, texto, 1, 12)
            base.fila_espaciadora(ws, fila); fila += 1

    # ── 9. Registro fotográfico complementario ──────────────────────────────
    sobrantes = [e for e in evidencias if e.id not in usadas]
    if sobrantes:
        secciones = []
        for e in sobrantes:
            if e.seccion not in secciones:
                secciones.append(e.seccion)
        fila_banner(ws, fila, numeros.titulo("REGISTRO FOTOGRÁFICO COMPLEMENTARIO"),
                    FILL_BURGUNDY, FONT_HEADER_WHITE)
        fila += 1
        base.fila_espaciadora(ws, fila); fila += 1
        for seccion in secciones:
            fotos = [e for e in sobrantes if e.seccion == seccion]
            fila = sub_titulo(ws, fila, seccion.upper())
            fila = _bloque_fotos_completo(ws, fila, fotos)

    # ── Cierre: certificados, firmas y fotos de técnicos ─────────────────────
    # Imágenes que no son evidencias: sin este cierre nunca llegaban al
    # reporte aunque estuvieran guardadas en la base.
    fila = base.bloque_imagenes_adjuntas(
        ws, fila, actividad,
        titulo=lambda: numeros.titulo("CERTIFICADOS, TÉCNICOS Y FIRMAS"))

    estilizar_fila(ws, 1, border=BORDER_NEGRO)
    return wb


def generar_excel_capex(actividad: SmuActividad) -> Workbook:
    """Devuelve el libro Excel del correctivo CAPEX (tipología que corresponda)."""
    return construir_libro(actividad)
